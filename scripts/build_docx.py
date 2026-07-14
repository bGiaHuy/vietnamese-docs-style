from docx import Document
from docx.shared import Pt, Cm, Inches, RGBColor, Emu
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn, nsdecls
from docx.oxml import parse_xml

BLACK = RGBColor(0x00, 0x00, 0x00)

def setup_document(doc, profile="administrative"):
    """Cấu hình trang giấy và style cơ bản theo profile."""
    section = doc.sections[0]
    section.page_width = Cm(21.0)   # A4
    section.page_height = Cm(29.7)  # A4
    section.top_margin = Cm(2.0)
    section.bottom_margin = Cm(2.0)
    
    # NĐ30: Lề trái 3.0-3.5cm, phải 1.5-2.0cm
    section.left_margin = Cm(3.0)
    section.right_margin = Cm(1.5)
    
    setup_styles(doc)

def setup_styles(doc):
    """Cấu hình các style mặc định."""
    # Normal (Body)
    style = doc.styles['Normal']
    sf = style.font
    sf.name = 'Times New Roman'
    sf.size = Pt(13) # NĐ30: 13-14pt
    sf.color.rgb = BLACK
    rPr = style.element.get_or_add_rPr()
    rPr.append(parse_xml(f'<w:rFonts {nsdecls("w")} w:eastAsia="Times New Roman"/>'))
    pf = style.paragraph_format
    pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    pf.space_before = Pt(0)
    pf.space_after = Pt(12)
    pf.line_spacing = 1.15

    # Heading 1
    h1 = doc.styles['Heading 1']
    h1.font.name = 'Times New Roman'
    h1.font.size = Pt(16)
    h1.font.bold = True
    h1.font.color.rgb = BLACK
    h1.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h1.paragraph_format.space_before = Pt(24)
    h1.paragraph_format.space_after = Pt(12)
    h1.paragraph_format.keep_with_next = True

    # Heading 2
    h2 = doc.styles['Heading 2']
    h2.font.name = 'Times New Roman'
    h2.font.size = Pt(13)
    h2.font.bold = True
    h2.font.color.rgb = BLACK
    h2.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h2.paragraph_format.space_before = Pt(18)
    h2.paragraph_format.space_after = Pt(8)
    h2.paragraph_format.keep_with_next = True

    # Heading 3
    h3 = doc.styles['Heading 3']
    h3.font.name = 'Times New Roman'
    h3.font.size = Pt(13)
    h3.font.bold = True
    h3.font.color.rgb = BLACK
    h3.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    h3.paragraph_format.keep_with_next = True

def add_body(doc, text):
    p = doc.add_paragraph(text)
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.color.rgb = BLACK
    return p

def add_blank(doc):
    p = doc.add_paragraph()
    pf = p.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    return p

def add_centered(doc, text, size=16, bold=True, italic=False):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.font.color.rgb = BLACK
    r.bold = bold
    r.italic = italic
    return p

def add_right(doc, text, bold=False, italic=False, size=13):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = p.add_run(text)
    r.font.name = 'Times New Roman'
    r.font.size = Pt(size)
    r.font.color.rgb = BLACK
    r.bold = bold
    r.italic = italic
    return p

def add_italic(doc, text):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    r = p.add_run(text)
    r.italic = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(13)
    r.font.color.rgb = BLACK
    return p

def add_bullet(doc, text, level=1, profile="administrative"):
    if profile in ["administrative", "minutes-administrative"]:
        prefixes = {1: "- ", 2: "+ ", 3: "* "}
        prefix = prefixes.get(level, "- ")
        p = doc.add_paragraph(prefix + text)
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.color.rgb = BLACK
        if level == 2:
            p.paragraph_format.left_indent = Cm(1.27)
        elif level == 3:
            p.paragraph_format.left_indent = Cm(2.54)
    else:
        style = 'List Bullet' if level == 1 else f'List Bullet {level}'
        p = doc.add_paragraph(text, style=style)
        for r in p.runs:
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.color.rgb = BLACK
    return p

def add_bullets(doc, items, level=1, profile="administrative"):
    for item in items:
        add_bullet(doc, item, level, profile)

def set_table_grid_style(table, profile="administrative"):
    table.style = 'Table Grid'
    if profile in ["administrative", "minutes-administrative"]:
        for row in table.rows:
            for cell in row.cells:
                tc = cell._tc
                tcPr = tc.get_or_add_tcPr()
                shading = tcPr.find(qn('w:shd'))
                if shading is not None:
                    tcPr.remove(shading)

def add_table(doc, headers, rows, profile="administrative"):
    tbl = doc.add_table(rows=1 + len(rows), cols=len(headers))
    set_table_grid_style(tbl, profile)
    for i, h in enumerate(headers):
        c = tbl.cell(0, i)
        c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = c.paragraphs[0].add_run(h)
        r.bold = True
        r.font.name = 'Times New Roman'
        r.font.size = Pt(13)
        r.font.color.rgb = BLACK
    for ri, row in enumerate(rows):
        for ci, val in enumerate(row):
            c = tbl.cell(ri + 1, ci)
            c.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = c.paragraphs[0].add_run(str(val))
            r.font.name = 'Times New Roman'
            r.font.size = Pt(13)
            r.font.color.rgb = BLACK
    return tbl

def add_cover_page(doc, org_name, title, subtitle, team_data, location="Hà Nội", year="2026"):
    add_centered(doc, org_name, 16, True)
    add_blank(doc)
    add_blank(doc)
    add_centered(doc, title, 16, True)
    add_blank(doc)
    if subtitle:
        add_centered(doc, subtitle, 14, True)
    add_blank(doc)
    if team_data:
        tbl = doc.add_table(rows=len(team_data), cols=2)
        tbl.style = 'Table Grid'
        tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
        for i, (lbl, val) in enumerate(team_data):
            c0 = tbl.cell(i, 0)
            c0.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
            r0 = c0.paragraphs[0].add_run(lbl)
            r0.bold = True; r0.font.name = 'Times New Roman'; r0.font.size = Pt(13)
            c1 = tbl.cell(i, 1)
            c1.paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.LEFT
            r1 = c1.paragraphs[0].add_run(val)
            r1.font.name = 'Times New Roman'; r1.font.size = Pt(13)
        for row in tbl.rows:
            row.cells[0].width = Cm(4)
            row.cells[1].width = Cm(10)
    add_blank(doc)
    add_centered(doc, f"{location}, {year}", 13, False, True)
    doc.add_page_break()

def add_cover_letter(doc, greeting, body_paragraphs, signer_name):
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p.add_run("THƯ NGỎ")
    r.bold = True
    r.font.name = 'Times New Roman'
    r.font.size = Pt(14)
    r.font.color.rgb = BLACK
    add_blank(doc)
    p_greet = doc.add_paragraph(greeting)
    p_greet.alignment = WD_ALIGN_PARAGRAPH.LEFT
    for run in p_greet.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(13)
        run.font.color.rgb = BLACK
    add_blank(doc)
    for bp_text in body_paragraphs:
        add_body(doc, bp_text)
        add_blank(doc)
    add_right(doc, "Trân trọng,", False, False)
    add_right(doc, signer_name, True)
    doc.add_page_break()

def add_h1(doc, text):
    p = doc.add_heading(text, level=1)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.color.rgb = BLACK
    return p

def add_h2(doc, text):
    p = doc.add_heading(text, level=2)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.color.rgb = BLACK
    return p

def add_h3(doc, text):
    p = doc.add_heading(text, level=3)
    for r in p.runs:
        r.font.name = 'Times New Roman'
        r.font.color.rgb = BLACK
    return p


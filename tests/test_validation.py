import os
import sys
import pytest
from docx import Document
from docx.shared import Cm, Inches, RGBColor, Pt

# Add scripts directory to path to import modules
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'scripts')))
from build_docx import setup_document, add_body, add_bullet, add_table, add_cover_page
from validate_docx import validate_doc

FIXTURES_DIR = os.path.join(os.path.dirname(__file__), 'fixtures')

def setup_module():
    if not os.path.exists(FIXTURES_DIR):
        os.makedirs(FIXTURES_DIR)

def test_valid_administrative_doc():
    filepath = os.path.join(FIXTURES_DIR, "test_administrative_valid.docx")
    doc = Document()
    setup_document(doc, profile="administrative")
    add_body(doc, "Đây là nội dung văn bản hành chính hợp lệ.")
    add_bullet(doc, "Mục 1", level=1, profile="administrative")
    
    doc.save(filepath)
    
    assert validate_doc(filepath, profile="administrative") == True

def test_invalid_page_size():
    filepath = os.path.join(FIXTURES_DIR, "test_invalid_size.docx")
    doc = Document()
    section = doc.sections[0]
    # Letter size
    section.page_width = Inches(8.5)
    section.page_height = Inches(11.0)
    
    doc.save(filepath)
    assert validate_doc(filepath, profile="administrative") == False

def test_invalid_margins():
    filepath = os.path.join(FIXTURES_DIR, "test_invalid_margins.docx")
    doc = Document()
    setup_document(doc, profile="administrative")
    
    # Overwrite with wrong margin
    section = doc.sections[0]
    section.right_margin = Cm(3.0) 
    
    doc.save(filepath)
    assert validate_doc(filepath, profile="administrative") == False

def test_invalid_bullets():
    filepath = os.path.join(FIXTURES_DIR, "test_invalid_bullets.docx")
    doc = Document()
    setup_document(doc, profile="administrative")
    
    # Add bullet dot
    doc.add_paragraph("• Mục không hợp lệ")
    
    doc.save(filepath)
    assert validate_doc(filepath, profile="administrative") == False

def test_invalid_color():
    filepath = os.path.join(FIXTURES_DIR, "test_invalid_color.docx")
    doc = Document()
    setup_document(doc, profile="administrative")
    
    p = doc.add_paragraph()
    r = p.add_run("Text màu đỏ")
    r.font.color.rgb = RGBColor(255, 0, 0)
    
    doc.save(filepath)
    assert validate_doc(filepath, profile="administrative") == False

def test_academic_allows_bullets():
    filepath = os.path.join(FIXTURES_DIR, "test_academic_bullets.docx")
    doc = Document()
    setup_document(doc, profile="academic")
    
    doc.add_paragraph("• Mục hợp lệ trong báo cáo học thuật")
    
    doc.save(filepath)
    assert validate_doc(filepath, profile="academic") == True
